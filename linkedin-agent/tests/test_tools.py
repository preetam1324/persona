"""Tests for all 7 LinkedIn agent tools."""

import os
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock

import pytest

from providers.mock import MockLinkedInProvider
from tools.campaign_manager import CampaignManagerTool
from tools.check_replies import CheckRepliesTool
from tools.converse import ConverseTool
from tools.export_leads import ExportLeadsTool
from tools.get_profiles import GetProfilesTool
from tools.search_leads import SearchLeadsTool
from tools.send_message import SendMessageTool


# ─── SearchLeadsTool ─────────────────────────────────────────────────────────


class TestSearchLeadsTool:
    @pytest.mark.asyncio
    async def test_search_returns_profiles(self, mock_provider):
        tool = SearchLeadsTool()
        tool.linkedin_provider = mock_provider
        tool.db = None
        result = await tool.execute(title="VP", max_results=5)
        assert result.success
        assert result.output["count"] >= 1
        assert "profiles" in result.output

    @pytest.mark.asyncio
    async def test_search_no_provider_returns_error(self):
        tool = SearchLeadsTool()
        tool.linkedin_provider = None
        tool.db = None
        result = await tool.execute(title="CTO")
        assert not result.success
        assert "not configured" in result.error

    @pytest.mark.asyncio
    async def test_search_stores_in_db(self, mock_provider, seeded_db):
        db, campaign_id = seeded_db
        tool = SearchLeadsTool()
        tool.linkedin_provider = mock_provider
        tool.db = db
        result = await tool.execute(
            title="Director", campaign_id=campaign_id
        )
        assert result.success
        assert result.output["campaign_id"] == campaign_id

    @pytest.mark.asyncio
    async def test_search_no_db_still_works(self, mock_provider):
        tool = SearchLeadsTool()
        tool.linkedin_provider = mock_provider
        tool.db = None
        result = await tool.execute(industry="Technology")
        assert result.success
        assert result.output["count"] >= 1

    @pytest.mark.asyncio
    async def test_search_empty_results(self, mock_provider):
        tool = SearchLeadsTool()
        tool.linkedin_provider = mock_provider
        tool.db = None
        result = await tool.execute(title="Nonexistent Title XYZ")
        assert result.success
        assert result.output["count"] == 0


# ─── GetProfilesTool ─────────────────────────────────────────────────────────


class TestGetProfilesTool:
    @pytest.mark.asyncio
    async def test_get_by_ids(self, mock_provider):
        tool = GetProfilesTool()
        tool.linkedin_provider = mock_provider
        tool.db = None
        result = await tool.execute(profile_ids=["mock-001", "mock-002"])
        assert result.success
        assert result.output["count"] == 2

    @pytest.mark.asyncio
    async def test_get_by_campaign_id(self, mock_provider, seeded_db):
        db, campaign_id = seeded_db
        tool = GetProfilesTool()
        tool.linkedin_provider = mock_provider
        tool.db = db
        result = await tool.execute(campaign_id=campaign_id)
        assert result.success
        assert result.output["count"] == 2

    @pytest.mark.asyncio
    async def test_get_no_provider_error(self):
        tool = GetProfilesTool()
        tool.linkedin_provider = None
        tool.db = None
        result = await tool.execute(profile_ids=["mock-001"])
        assert not result.success

    @pytest.mark.asyncio
    async def test_get_empty_list(self, mock_provider):
        tool = GetProfilesTool()
        tool.linkedin_provider = mock_provider
        tool.db = None
        result = await tool.execute()
        assert result.success
        assert result.output["count"] == 0


# ─── ExportLeadsTool ─────────────────────────────────────────────────────────


class TestExportLeadsTool:
    @pytest.mark.asyncio
    async def test_export_generates_xlsx(self, seeded_db, tmp_path, monkeypatch):
        db, campaign_id = seeded_db
        # Redirect exports to tmp_path
        monkeypatch.chdir(tmp_path)
        tool = ExportLeadsTool()
        tool.db = db
        result = await tool.execute(campaign_id=campaign_id)
        assert result.success
        assert result.output["count"] == 2
        assert result.output["filename"].endswith(".xlsx")
        assert "/v1/agent/downloads/" in result.output["download_url"]

        # Verify the file actually exists
        filepath = tmp_path / "data" / "exports" / result.output["filename"]
        assert filepath.exists()

    @pytest.mark.asyncio
    async def test_export_no_leads(self, test_db, tmp_path, monkeypatch):
        from db.repository import CampaignRepository

        async with test_db.session() as session:
            repo = CampaignRepository(session)
            campaign = await repo.create(name="Empty", criteria={})
            campaign_id = campaign.id

        monkeypatch.chdir(tmp_path)
        tool = ExportLeadsTool()
        tool.db = test_db
        result = await tool.execute(campaign_id=campaign_id)
        assert result.success
        assert result.output["count"] == 0

    @pytest.mark.asyncio
    async def test_export_no_db_error(self):
        tool = ExportLeadsTool()
        tool.db = None
        result = await tool.execute(campaign_id="some-id")
        assert not result.success
        assert "Database not configured" in result.error

    @pytest.mark.asyncio
    async def test_export_with_status_filter(self, seeded_db, tmp_path, monkeypatch):
        db, campaign_id = seeded_db
        monkeypatch.chdir(tmp_path)
        tool = ExportLeadsTool()
        tool.db = db
        # All leads are "new", filter by "contacted" should find 0
        result = await tool.execute(
            campaign_id=campaign_id, status_filter="contacted"
        )
        assert result.success
        assert result.output["count"] == 0

    @pytest.mark.asyncio
    async def test_export_xlsx_columns(self, seeded_db, tmp_path, monkeypatch):
        db, campaign_id = seeded_db
        monkeypatch.chdir(tmp_path)
        tool = ExportLeadsTool()
        tool.db = db
        result = await tool.execute(campaign_id=campaign_id)

        from openpyxl import load_workbook

        filepath = tmp_path / "data" / "exports" / result.output["filename"]
        wb = load_workbook(filepath)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        assert headers == [
            "Name", "Title", "Company", "Location",
            "Headline", "Profile URL", "Status",
        ]
        assert ws.max_row == 3  # header + 2 leads


# ─── SendMessageTool ─────────────────────────────────────────────────────────


class TestSendMessageTool:
    @pytest.mark.asyncio
    async def test_send_template_to_all_new(self, mock_provider, seeded_db):
        db, campaign_id = seeded_db
        tool = SendMessageTool()
        tool.linkedin_provider = mock_provider
        tool.db = db
        tool.llm_provider = None
        result = await tool.execute(
            campaign_id=campaign_id,
            message_template="Hi {name}, great work at {company}!",
        )
        assert result.success
        assert result.output["sent"] == 2
        assert result.output["failed"] == 0
        assert result.output["personalized"] is False

    @pytest.mark.asyncio
    async def test_send_to_specific_ids(self, mock_provider, seeded_db):
        db, campaign_id = seeded_db
        tool = SendMessageTool()
        tool.linkedin_provider = mock_provider
        tool.db = db
        tool.llm_provider = None
        result = await tool.execute(
            campaign_id=campaign_id,
            message_template="Hi {name}!",
            lead_ids=["mock-001"],
        )
        assert result.success
        assert result.output["sent"] == 1

    @pytest.mark.asyncio
    async def test_template_substitution(self, mock_provider, seeded_db):
        db, campaign_id = seeded_db
        tool = SendMessageTool()
        tool.linkedin_provider = mock_provider
        tool.db = db
        tool.llm_provider = None
        result = await tool.execute(
            campaign_id=campaign_id,
            message_template="Hi {name}, your role as {title} at {company} is impressive!",
            lead_ids=["mock-001"],
        )
        assert result.success
        assert result.output["sent"] == 1

    @pytest.mark.asyncio
    async def test_send_updates_status_to_contacted(self, mock_provider, seeded_db):
        db, campaign_id = seeded_db
        tool = SendMessageTool()
        tool.linkedin_provider = mock_provider
        tool.db = db
        tool.llm_provider = None
        await tool.execute(
            campaign_id=campaign_id,
            message_template="Hi {name}!",
        )
        # Verify status changed
        from db.models import LeadStatus
        from db.repository import LeadRepository

        async with db.session() as session:
            repo = LeadRepository(session)
            contacted = await repo.get_campaign_leads(
                campaign_id, status=LeadStatus.contacted
            )
            assert len(contacted) == 2

    @pytest.mark.asyncio
    async def test_send_records_message_in_db(self, mock_provider, seeded_db):
        db, campaign_id = seeded_db
        tool = SendMessageTool()
        tool.linkedin_provider = mock_provider
        tool.db = db
        tool.llm_provider = None
        await tool.execute(
            campaign_id=campaign_id,
            message_template="Hi {name}!",
            lead_ids=["mock-001"],
        )
        from db.models import Message, MessageDirection
        from sqlalchemy import select

        async with db.session() as session:
            result = await session.execute(select(Message))
            messages = list(result.scalars().all())
            assert len(messages) >= 1
            assert messages[0].direction == MessageDirection.outbound

    @pytest.mark.asyncio
    async def test_send_no_provider_error(self, seeded_db):
        db, campaign_id = seeded_db
        tool = SendMessageTool()
        tool.linkedin_provider = None
        tool.db = db
        tool.llm_provider = None
        result = await tool.execute(
            campaign_id=campaign_id,
            message_template="Hi!",
        )
        assert not result.success

    @pytest.mark.asyncio
    async def test_send_no_leads_to_message(self, mock_provider, seeded_db):
        db, campaign_id = seeded_db
        tool = SendMessageTool()
        tool.linkedin_provider = mock_provider
        tool.db = db
        tool.llm_provider = None
        result = await tool.execute(
            campaign_id=campaign_id,
            message_template="Hi!",
            lead_ids=["nonexistent-id"],
        )
        assert result.success
        assert result.output["sent"] == 0

    @pytest.mark.asyncio
    async def test_personalize_flag(self, mock_provider, seeded_db):
        """With personalize=True but no LLM provider, falls back to template."""
        db, campaign_id = seeded_db
        tool = SendMessageTool()
        tool.linkedin_provider = mock_provider
        tool.db = db
        tool.llm_provider = None  # No LLM — should fall back gracefully
        result = await tool.execute(
            campaign_id=campaign_id,
            message_template="Hi {name}!",
            personalize=True,
            lead_ids=["mock-001"],
        )
        assert result.success
        assert result.output["sent"] == 1

    @pytest.mark.asyncio
    async def test_personalize_with_mock_llm(self, mock_provider, seeded_db):
        """With a mocked LLM, personalize should produce output."""
        db, campaign_id = seeded_db

        # Mock LLM provider
        mock_llm = MagicMock()
        mock_response = MagicMock()
        mock_response.content = "Hey Alice, your cloud infra work is amazing!"
        mock_llm.complete = AsyncMock(return_value=mock_response)

        tool = SendMessageTool()
        tool.linkedin_provider = mock_provider
        tool.db = db
        tool.llm_provider = mock_llm
        result = await tool.execute(
            campaign_id=campaign_id,
            message_template="Hi {name}!",
            personalize=True,
            lead_ids=["mock-001"],
        )
        assert result.success
        assert result.output["personalized"] is True
        mock_llm.complete.assert_called_once()


# ─── CheckRepliesTool ─────────────────────────────────────────────────────────


class TestCheckRepliesTool:
    @pytest.mark.asyncio
    async def test_no_replies(self, mock_provider, seeded_db):
        db, campaign_id = seeded_db
        tool = CheckRepliesTool()
        tool.linkedin_provider = mock_provider
        tool.db = db
        result = await tool.execute(campaign_id=campaign_id)
        assert result.success
        assert result.output["reply_count"] == 0

    @pytest.mark.asyncio
    async def test_detects_reply(self, mock_provider, seeded_db):
        db, campaign_id = seeded_db
        # First send a message so mock provider simulates a reply
        await mock_provider.send_message("mock-001", "Hello!")

        # Mark lead as contacted in DB
        from db.models import LeadStatus
        from db.repository import LeadRepository

        async with db.session() as session:
            repo = LeadRepository(session)
            pairs = await repo.get_campaign_leads(campaign_id)
            for cl, lead in pairs:
                if lead.linkedin_profile_id == "mock-001":
                    await repo.update_lead_status(cl.id, LeadStatus.contacted)

        tool = CheckRepliesTool()
        tool.linkedin_provider = mock_provider
        tool.db = db
        result = await tool.execute(campaign_id=campaign_id)
        assert result.success
        assert result.output["reply_count"] == 1
        assert result.output["replies"][0]["lead_name"] == "Alice Johnson"

    @pytest.mark.asyncio
    async def test_updates_status_to_replied(self, mock_provider, seeded_db):
        db, campaign_id = seeded_db
        await mock_provider.send_message("mock-001", "Hello!")

        from db.models import LeadStatus
        from db.repository import LeadRepository

        async with db.session() as session:
            repo = LeadRepository(session)
            pairs = await repo.get_campaign_leads(campaign_id)
            for cl, lead in pairs:
                if lead.linkedin_profile_id == "mock-001":
                    await repo.update_lead_status(cl.id, LeadStatus.contacted)

        tool = CheckRepliesTool()
        tool.linkedin_provider = mock_provider
        tool.db = db
        await tool.execute(campaign_id=campaign_id)

        async with db.session() as session:
            repo = LeadRepository(session)
            replied = await repo.get_campaign_leads(
                campaign_id, status=LeadStatus.replied
            )
            assert len(replied) == 1

    @pytest.mark.asyncio
    async def test_no_provider_error(self, seeded_db):
        db, campaign_id = seeded_db
        tool = CheckRepliesTool()
        tool.linkedin_provider = None
        tool.db = db
        result = await tool.execute(campaign_id=campaign_id)
        assert not result.success

    @pytest.mark.asyncio
    async def test_no_db_error(self, mock_provider):
        tool = CheckRepliesTool()
        tool.linkedin_provider = mock_provider
        tool.db = None
        result = await tool.execute()
        assert not result.success

    @pytest.mark.asyncio
    async def test_check_all_active_campaigns(self, mock_provider, seeded_db):
        db, campaign_id = seeded_db
        tool = CheckRepliesTool()
        tool.linkedin_provider = mock_provider
        tool.db = db
        # No campaign_id — should check all active
        result = await tool.execute()
        assert result.success


# ─── CampaignManagerTool ──────────────────────────────────────────────────────


class TestCampaignManagerTool:
    @pytest.mark.asyncio
    async def test_create_campaign(self, test_db):
        tool = CampaignManagerTool()
        tool.db = test_db
        tool.scheduler = None
        result = await tool.execute(
            action="create",
            name="Q3 Leads",
            criteria={"title": "CTO"},
            interval_weeks=3,
        )
        assert result.success
        assert result.output["name"] == "Q3 Leads"
        assert result.output["campaign_id"] is not None

    @pytest.mark.asyncio
    async def test_list_campaigns(self, seeded_db):
        db, _ = seeded_db
        tool = CampaignManagerTool()
        tool.db = db
        tool.scheduler = None
        result = await tool.execute(action="list")
        assert result.success
        assert result.output["count"] >= 1

    @pytest.mark.asyncio
    async def test_campaign_status(self, seeded_db):
        db, campaign_id = seeded_db
        tool = CampaignManagerTool()
        tool.db = db
        tool.scheduler = None
        result = await tool.execute(action="status", campaign_id=campaign_id)
        assert result.success
        assert result.output["name"] == "Test Campaign"
        assert result.output["total_leads"] == 2

    @pytest.mark.asyncio
    async def test_pause_campaign(self, seeded_db):
        db, campaign_id = seeded_db
        tool = CampaignManagerTool()
        tool.db = db
        tool.scheduler = None
        result = await tool.execute(action="pause", campaign_id=campaign_id)
        assert result.success
        assert result.output["status"] == "paused"

    @pytest.mark.asyncio
    async def test_resume_campaign(self, seeded_db):
        db, campaign_id = seeded_db
        tool = CampaignManagerTool()
        tool.db = db
        tool.scheduler = None
        await tool.execute(action="pause", campaign_id=campaign_id)
        result = await tool.execute(action="resume", campaign_id=campaign_id)
        assert result.success
        assert result.output["status"] == "active"

    @pytest.mark.asyncio
    async def test_delete_campaign(self, seeded_db):
        db, campaign_id = seeded_db
        tool = CampaignManagerTool()
        tool.db = db
        tool.scheduler = None
        result = await tool.execute(action="delete", campaign_id=campaign_id)
        assert result.success

    @pytest.mark.asyncio
    async def test_no_db_error(self):
        tool = CampaignManagerTool()
        tool.db = None
        tool.scheduler = None
        result = await tool.execute(action="list")
        assert not result.success

    @pytest.mark.asyncio
    async def test_missing_campaign_id(self, test_db):
        tool = CampaignManagerTool()
        tool.db = test_db
        tool.scheduler = None
        result = await tool.execute(action="pause")
        assert not result.success
        assert "campaign_id is required" in result.error

    @pytest.mark.asyncio
    async def test_nonexistent_campaign_status(self, test_db):
        tool = CampaignManagerTool()
        tool.db = test_db
        tool.scheduler = None
        result = await tool.execute(
            action="status", campaign_id="nonexistent"
        )
        assert not result.success
        assert "not found" in result.error


# ─── ConverseTool ────────────────────────────────────────────────────────────


class TestConverseTool:
    @pytest.mark.asyncio
    async def test_send_message(self, mock_provider):
        tool = ConverseTool()
        tool.linkedin_provider = mock_provider
        tool.db = None
        tool.llm_provider = None
        result = await tool.execute(
            conversation_id="conv-001",
            message="Thanks for your interest!",
        )
        assert result.success
        assert result.output["status"] == "sent"

    @pytest.mark.asyncio
    async def test_no_message_no_auto_reply_error(self, mock_provider):
        tool = ConverseTool()
        tool.linkedin_provider = mock_provider
        tool.db = None
        tool.llm_provider = None
        result = await tool.execute(
            conversation_id="conv-001",
            auto_reply=False,
        )
        assert not result.success
        assert "No message" in result.error

    @pytest.mark.asyncio
    async def test_auto_reply_with_mock_llm(self, mock_provider):
        # Must have a conversation to generate reply from
        await mock_provider.send_message("mock-001", "Hello!")

        mock_llm = MagicMock()
        mock_response = MagicMock()
        mock_response.content = "Great to hear from you! Let's schedule a call."
        mock_llm.complete = AsyncMock(return_value=mock_response)

        tool = ConverseTool()
        tool.linkedin_provider = mock_provider
        tool.db = None
        tool.llm_provider = mock_llm
        result = await tool.execute(
            conversation_id="conv-001",
            auto_reply=True,
        )
        assert result.success
        assert "schedule a call" in result.output["message_sent"]

    @pytest.mark.asyncio
    async def test_no_provider_error(self):
        tool = ConverseTool()
        tool.linkedin_provider = None
        tool.db = None
        tool.llm_provider = None
        result = await tool.execute(
            conversation_id="conv-001",
            message="Hello",
        )
        assert not result.success

    @pytest.mark.asyncio
    async def test_records_in_db(self, mock_provider, seeded_db):
        db, campaign_id = seeded_db
        # Create a conversation record in DB first
        from db.repository import LeadRepository, MessageRepository

        async with db.session() as session:
            lead_repo = LeadRepository(session)
            pairs = await lead_repo.get_campaign_leads(campaign_id)
            cl_id = pairs[0][0].id
            msg_repo = MessageRepository(session)
            await msg_repo.upsert_conversation(
                campaign_lead_id=cl_id,
                linkedin_conversation_id="conv-001",
            )

        tool = ConverseTool()
        tool.linkedin_provider = mock_provider
        tool.db = db
        tool.llm_provider = None
        result = await tool.execute(
            conversation_id="conv-001",
            message="Follow up message",
        )
        assert result.success

        # Verify message recorded
        from db.models import Message
        from sqlalchemy import select

        async with db.session() as session:
            msgs = await session.execute(select(Message))
            all_msgs = list(msgs.scalars().all())
            assert any("Follow up" in m.content for m in all_msgs)

    @pytest.mark.asyncio
    async def test_auto_reply_conversation_not_found(self, mock_provider):
        mock_llm = MagicMock()
        mock_llm.complete = AsyncMock()

        tool = ConverseTool()
        tool.linkedin_provider = mock_provider
        tool.db = None
        tool.llm_provider = mock_llm
        result = await tool.execute(
            conversation_id="nonexistent-conv",
            auto_reply=True,
        )
        assert not result.success
